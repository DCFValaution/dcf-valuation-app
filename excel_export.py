"""
Builds a downloadable .xlsx DCF model mirroring AAPL_DCF_Model.xlsx.

EVERYTHING IS A FORMULA
-----------------------
The point of exporting a workbook rather than a PDF is that the recipient can
argue with it. So only two kinds of cell hold literal numbers:

  * the assumption block (blue on yellow) - the things a user changes
  * the base-year actuals (blue) - what the company reported

Every other figure, including all 25 sensitivity cells, is an Excel formula
referencing those inputs. Change WACC in B11 and the projection, terminal
value, equity bridge, and the whole sensitivity grid recalculate - including
the grid's own axis labels, which are anchored to the WACC and terminal growth
cells so the table stays centred on whatever the user sets.

This is why the module does not simply write the numbers that analysis.py
already computed: a workbook of hardcoded values would look identical on
opening and be inert thereafter.

CACHED VALUES
-------------
openpyxl writes formulas without cached results, so a freshly generated file
has no stored values. `fullCalcOnLoad` is set so Excel (and LibreOffice, and
Google Sheets) computes everything on open. The consequence worth knowing: a
tool that reads the file without an evaluation engine sees formula strings,
not numbers. That is expected, not corruption.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analysis import ValuationReport, honesty_note

FONT_NAME = "Arial"

# Industry-standard colour conventions: blue = hardcoded input, black = formula.
BLUE = "0000FF"
BLACK = "000000"
GREY = "595959"
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
HEADER_FILL = PatternFill("solid", start_color="D9E1F2")
TITLE_FILL = PatternFill("solid", start_color="1F3864")

CURRENCY = '$#,##0;($#,##0);"-"'
PER_SHARE = '$#,##0.00;($#,##0.00);"-"'
PERCENT_2DP = '0.00%'
PERCENT_1DP = '0.0%'
SHARE_COUNT = '#,##0.0'
FACTOR = '0.000'

THIN_TOP = Border(top=Side(style="thin", color="808080"))

# --- Fixed row layout. Only columns vary with the projection horizon. -------
R_TITLE = 1
R_SUBTITLE = 2
R_ASSUMPTIONS_HEAD = 4
R_GROWTH = 5
R_MARGIN = 6
R_TAX = 7
R_DA = 8
R_CAPEX = 9
R_NWC = 10
R_WACC = 11
R_TERMINAL = 12
R_YEARS = 13
R_BASE_HEAD = 15
R_REVENUE_IN = 16
R_DEBT_IN = 17
R_CASH_IN = 18
R_SHARES_IN = 19
R_PRICE_IN = 20
R_PROJ_HEAD = 22
R_COLUMN_HEADS = 23
R_PERIOD = 24
R_REVENUE = 25
R_GROWTH_PCT = 26
R_EBIT = 27
R_TAXES = 28
R_NOPAT = 29
R_DA_ADD = 30
R_CAPEX_LESS = 31
R_NWC_LESS = 32
R_UFCF = 33
R_DISCOUNT = 34
R_PV_UFCF = 35
R_EV_HEAD = 37
R_PV_SUM = 38
R_TV = 39
R_PV_TV = 40
R_EV = 41
R_PLUS_CASH = 42
R_LESS_DEBT = 43
R_EQUITY = 44
R_SHARES = 45
R_INTRINSIC = 46
R_PRICE = 47
R_UPSIDE = 48
R_TV_PCT = 49
R_SENS_HEAD = 51
R_SENS_LEGEND = 52
R_SENS_COLS = 53
R_SENS_FIRST = 54
R_SENS_LAST = 58
R_NOTES_HEAD = 60

SENSITIVITY_OFFSETS = (-0.01, -0.005, 0.0, 0.005, 0.01)


@dataclass
class _Layout:
    """Column geometry, which depends on the projection horizon."""
    years: int

    @property
    def base_col(self) -> int:
        return 2  # B

    @property
    def first_col(self) -> int:
        return 3  # C

    @property
    def last_col(self) -> int:
        return 2 + self.years

    @property
    def first_letter(self) -> str:
        return get_column_letter(self.first_col)

    @property
    def last_letter(self) -> str:
        return get_column_letter(self.last_col)

    def letter(self, index: int) -> str:
        return get_column_letter(index)

    @property
    def projection_cols(self) -> range:
        return range(self.first_col, self.last_col + 1)


def _style(cell, *, bold=False, color=BLACK, size=10, fmt=None,
           fill=None, align=None, italic=False, wrap=False):
    cell.font = Font(name=FONT_NAME, bold=bold, color=color, size=size, italic=italic)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align or wrap:
        cell.alignment = Alignment(horizontal=align, wrap_text=wrap)
    return cell


def _label(ws, row, text, *, bold=False, indent=0, color=BLACK, size=10, italic=False):
    cell = ws.cell(row=row, column=1, value=("   " * indent) + text)
    return _style(cell, bold=bold, color=color, size=size, italic=italic)


def _input(ws, row, value, fmt, *, note=None, source=None):
    """A hardcoded, user-editable input: blue text on a highlighted fill."""
    cell = ws.cell(row=row, column=2, value=value)
    _style(cell, color=BLUE, fmt=fmt, fill=INPUT_FILL)
    if source:
        _style(ws.cell(row=row, column=3, value=source), color=GREY, size=9)
    if note:
        _style(ws.cell(row=row, column=4, value=note), color=GREY, size=9)
    return cell


def _formula(ws, row, col, formula, fmt, *, bold=False, top_border=False):
    cell = ws.cell(row=row, column=col, value=formula)
    _style(cell, bold=bold, fmt=fmt)
    if top_border:
        cell.border = THIN_TOP
    return cell


def _section(ws, row, text, layout: _Layout):
    cell = _label(ws, row, text, bold=True, size=11)
    cell.fill = HEADER_FILL
    for col in range(2, layout.last_col + 1):
        ws.cell(row=row, column=col).fill = HEADER_FILL


def build_workbook(report: ValuationReport) -> Workbook:
    """
    Build a live-formula DCF workbook for *report*.

    The caller is responsible for having checked suitability: this function
    assumes report.result is populated. Generating a workbook for a company
    the guard refused would hand the user exactly the misleading artefact the
    guard exists to prevent.
    """
    if getattr(report, "method", "dcf") != "dcf":
        raise ValueError(
            f"Refusing to build a workbook for {report.ticker}: the export builds "
            f"DCF models, and this company was valued with a {report.method.upper()}."
        )
    if report.result is None:
        raise ValueError(
            f"Refusing to build a workbook for {report.ticker}: the suitability "
            "guard produced no valuation."
        )

    a = report.derived.assumptions
    base = report.base
    layout = _Layout(years=a.projection_years)
    provenance = report.derived.provenance
    today = _dt.date.today().isoformat()

    try:
        base_year = int(report.fiscal_year)
    except (TypeError, ValueError):
        base_year = None

    wb = Workbook()
    ws = wb.active
    ws.title = "DCF Model"

    # Excel stores no cached results for openpyxl-written formulas; force a
    # full calculation when the file is opened.
    wb.calculation.fullCalcOnLoad = True

    # --- Title -------------------------------------------------------------
    title = _label(ws, R_TITLE,
                   f"{report.company_name}  ({report.ticker})  -  "
                   "Discounted Cash Flow (DCF) Valuation",
                   bold=True, size=14)
    title.font = Font(name=FONT_NAME, bold=True, size=14, color="FFFFFF")
    for col in range(1, layout.last_col + 1):
        ws.cell(row=R_TITLE, column=col).fill = TITLE_FILL

    fy_text = f"FY{report.fiscal_year}" if base_year else "the latest reported fiscal year"
    _label(ws, R_SUBTITLE,
           f"$ in millions except per-share data.  Base year: {fy_text}.  "
           "Every white cell is a formula - edit the shaded input cells and the "
           "whole model recalculates.",
           italic=True, color=GREY, size=9)

    # --- 1. Assumptions ----------------------------------------------------
    _section(ws, R_ASSUMPTIONS_HEAD,
             "1.  ASSUMPTIONS   (shaded blue-text cells are editable inputs)", layout)

    assumption_rows = [
        (R_GROWTH, "Revenue growth rate (per year)", "revenue_growth", PERCENT_2DP),
        (R_MARGIN, "Operating margin (EBIT % of revenue)", "operating_margin", PERCENT_2DP),
        (R_TAX, "Tax rate", "tax_rate", PERCENT_2DP),
        (R_DA, "Depreciation & amortisation (% of revenue)", "da_pct", PERCENT_2DP),
        (R_CAPEX, "Capital expenditures (% of revenue)", "capex_pct", PERCENT_2DP),
        (R_NWC, "Change in net working capital (% of revenue growth)", "nwc_pct", PERCENT_2DP),
        (R_WACC, "WACC (discount rate)", "wacc", PERCENT_2DP),
        (R_TERMINAL, "Terminal growth rate (g)", "terminal_growth", PERCENT_2DP),
    ]
    for row, text, name, fmt in assumption_rows:
        _label(ws, row, text)
        prov = provenance[name]
        _input(ws, row, getattr(a, name), fmt, source=prov.source, note=prov.detail)

    # Structural rather than editable: the projection grid is generated with a
    # fixed number of columns, so changing this cell in Excel would not add
    # years. Left black and unshaded to signal it is not an input.
    _label(ws, R_YEARS, "Projection period (years)")
    _style(ws.cell(row=R_YEARS, column=2, value=a.projection_years), fmt="0")
    _style(ws.cell(row=R_YEARS, column=3, value="structural"), color=GREY, size=9)
    _style(ws.cell(row=R_YEARS, column=4,
                   value="regenerate the file to change the horizon"), color=GREY, size=9)

    # --- 2. Base-year data -------------------------------------------------
    _section(ws, R_BASE_HEAD, f"2.  BASE-YEAR DATA & MARKET   ({fy_text} actuals)", layout)
    meta = report.meta
    source_note = f"Source: Financial Modeling Prep, {today}"

    base_rows = [
        (R_REVENUE_IN, f"Revenue - {fy_text} ($mm)", base.revenue, CURRENCY,
         meta.get("revenue_source", "")),
        (R_DEBT_IN, "Total debt ($mm)", base.total_debt, CURRENCY,
         meta.get("debt_source", "")),
        (R_CASH_IN, "Cash & investments ($mm)", base.cash, CURRENCY,
         meta.get("cash_source", "")),
        (R_SHARES_IN, "Shares outstanding - diluted (millions)", base.shares, SHARE_COUNT,
         meta.get("shares_source", "")),
        (R_PRICE_IN, "Current share price ($)", base.current_price, PER_SHARE,
         meta.get("price_source", "")),
    ]
    for row, text, value, fmt, field_source in base_rows:
        _label(ws, row, text)
        _input(ws, row, value, fmt, source=source_note, note=field_source)

    # --- 3. Projection -----------------------------------------------------
    _section(ws, R_PROJ_HEAD, "3.  UNLEVERED FREE CASH FLOW PROJECTION ($mm)", layout)

    _style(ws.cell(row=R_COLUMN_HEADS, column=1, value="Line item ($mm)"), bold=True)
    base_label = f"{fy_text} (Base)" if base_year else "Base"
    _style(ws.cell(row=R_COLUMN_HEADS, column=layout.base_col, value=base_label),
           bold=True, align="right")
    for offset, col in enumerate(layout.projection_cols, start=1):
        heading = f"FY{base_year + offset}E" if base_year else f"Year {offset}E"
        _style(ws.cell(row=R_COLUMN_HEADS, column=col, value=heading),
               bold=True, align="right")

    _label(ws, R_PERIOD, "Period (t)", italic=True, color=GREY)
    for offset, col in enumerate(layout.projection_cols, start=1):
        _style(ws.cell(row=R_PERIOD, column=col, value=offset), fmt="0", align="right")

    B = layout.letter(layout.base_col)

    # Revenue: base links to the input cell, then compounds by the growth rate.
    _label(ws, R_REVENUE, "Revenue")
    _formula(ws, R_REVENUE, layout.base_col, f"=$B${R_REVENUE_IN}", CURRENCY)
    _label(ws, R_GROWTH_PCT, "Revenue growth %", indent=1, italic=True, color=GREY)
    _label(ws, R_EBIT, "Operating income (EBIT)")
    _label(ws, R_TAXES, "Less: taxes on EBIT", indent=1)
    _label(ws, R_NOPAT, "NOPAT (EBIT after tax)")
    _label(ws, R_DA_ADD, "Plus: depreciation & amortisation", indent=1)
    _label(ws, R_CAPEX_LESS, "Less: capital expenditures", indent=1)
    _label(ws, R_NWC_LESS, "Less: change in net working capital", indent=1)
    _label(ws, R_UFCF, "Unlevered free cash flow", bold=True)
    _label(ws, R_DISCOUNT, "Discount factor @ WACC", indent=1, italic=True, color=GREY)
    _label(ws, R_PV_UFCF, "PV of unlevered FCF", bold=True)

    # Base-year column: the reported starting point, on the same formula basis
    # as the projection so the two are directly comparable.
    _formula(ws, R_EBIT, layout.base_col, f"={B}{R_REVENUE}*$B${R_MARGIN}", CURRENCY)
    _formula(ws, R_TAXES, layout.base_col, f"=-{B}{R_EBIT}*$B${R_TAX}", CURRENCY)
    _formula(ws, R_NOPAT, layout.base_col, f"={B}{R_EBIT}+{B}{R_TAXES}", CURRENCY)
    _formula(ws, R_DA_ADD, layout.base_col, f"={B}{R_REVENUE}*$B${R_DA}", CURRENCY)
    _formula(ws, R_CAPEX_LESS, layout.base_col, f"=-{B}{R_REVENUE}*$B${R_CAPEX}", CURRENCY)
    _style(ws.cell(row=R_NWC_LESS, column=layout.base_col, value=0), fmt=CURRENCY)
    _formula(ws, R_UFCF, layout.base_col,
             f"=SUM({B}{R_NOPAT}:{B}{R_NWC_LESS})", CURRENCY,
             bold=True, top_border=True)

    for col in layout.projection_cols:
        c = layout.letter(col)
        prev = layout.letter(col - 1)

        _formula(ws, R_REVENUE, col, f"={prev}{R_REVENUE}*(1+$B${R_GROWTH})", CURRENCY)
        _formula(ws, R_GROWTH_PCT, col, f"={c}{R_REVENUE}/{prev}{R_REVENUE}-1", PERCENT_1DP)
        _formula(ws, R_EBIT, col, f"={c}{R_REVENUE}*$B${R_MARGIN}", CURRENCY)
        _formula(ws, R_TAXES, col, f"=-{c}{R_EBIT}*$B${R_TAX}", CURRENCY)
        _formula(ws, R_NOPAT, col, f"={c}{R_EBIT}+{c}{R_TAXES}", CURRENCY)
        _formula(ws, R_DA_ADD, col, f"={c}{R_REVENUE}*$B${R_DA}", CURRENCY)
        _formula(ws, R_CAPEX_LESS, col, f"=-{c}{R_REVENUE}*$B${R_CAPEX}", CURRENCY)
        _formula(ws, R_NWC_LESS, col,
                 f"=-({c}{R_REVENUE}-{prev}{R_REVENUE})*$B${R_NWC}", CURRENCY)
        _formula(ws, R_UFCF, col, f"=SUM({c}{R_NOPAT}:{c}{R_NWC_LESS})", CURRENCY,
                 bold=True, top_border=True)
        _formula(ws, R_DISCOUNT, col, f"=1/(1+$B${R_WACC})^{c}${R_PERIOD}", FACTOR)
        _formula(ws, R_PV_UFCF, col, f"={c}{R_UFCF}*{c}{R_DISCOUNT}", CURRENCY, bold=True)

    first, last = layout.first_letter, layout.last_letter

    # --- 4. Enterprise & equity value --------------------------------------
    _section(ws, R_EV_HEAD, "4.  ENTERPRISE & EQUITY VALUE", layout)

    ev_rows = [
        (R_PV_SUM, "Sum of PV of explicit UFCF",
         f"=SUM({first}{R_PV_UFCF}:{last}{R_PV_UFCF})", CURRENCY, False),
        (R_TV, "Terminal value at end of forecast (Gordon growth)",
         f"={last}{R_UFCF}*(1+$B${R_TERMINAL})/($B${R_WACC}-$B${R_TERMINAL})",
         CURRENCY, False),
        (R_PV_TV, "PV of terminal value",
         f"=$B${R_TV}*{last}{R_DISCOUNT}", CURRENCY, False),
        (R_EV, "Enterprise value", f"=$B${R_PV_SUM}+$B${R_PV_TV}", CURRENCY, True),
        (R_PLUS_CASH, "Plus: cash & investments", f"=$B${R_CASH_IN}", CURRENCY, False),
        (R_LESS_DEBT, "Less: total debt", f"=-$B${R_DEBT_IN}", CURRENCY, False),
        (R_EQUITY, "Equity value",
         f"=$B${R_EV}+$B${R_PLUS_CASH}+$B${R_LESS_DEBT}", CURRENCY, True),
        (R_SHARES, "Shares outstanding - diluted (millions)",
         f"=$B${R_SHARES_IN}", SHARE_COUNT, False),
        (R_INTRINSIC, "INTRINSIC VALUE PER SHARE",
         f"=$B${R_EQUITY}/$B${R_SHARES}", PER_SHARE, True),
        (R_PRICE, "Current share price", f"=$B${R_PRICE_IN}", PER_SHARE, False),
        (R_UPSIDE, "Implied upside / (downside)",
         f"=$B${R_INTRINSIC}/$B${R_PRICE}-1", PERCENT_1DP, False),
        (R_TV_PCT, "Terminal value as % of enterprise value",
         f"=$B${R_PV_TV}/$B${R_EV}", PERCENT_1DP, False),
    ]
    for row, text, formula, fmt, bold in ev_rows:
        indent = 1 if text.startswith(("Plus:", "Less:")) else 0
        _label(ws, row, text, bold=bold, indent=indent)
        _formula(ws, row, 2, formula, fmt, bold=bold, top_border=bold)

    _style(ws.cell(row=R_INTRINSIC, column=1), bold=True, size=11)
    ws.cell(row=R_INTRINSIC, column=2).font = Font(name=FONT_NAME, bold=True, size=11)

    # --- 5. Sensitivity ----------------------------------------------------
    _section(ws, R_SENS_HEAD,
             "5.  SENSITIVITY - Intrinsic Value per Share ($)", layout)
    _label(ws, R_SENS_LEGEND,
           "Rows = WACC   |   Columns = terminal growth rate (g).  Axis labels are "
           "formulas anchored to B11/B12, so the grid re-centres when you change them.",
           italic=True, color=GREY, size=9)

    _style(ws.cell(row=R_SENS_COLS, column=2, value="WACC \\ g"),
           bold=True, align="center", fill=HEADER_FILL)

    # Axis headers are formulas offset from the assumption cells, so the table
    # follows whatever WACC and g the user sets rather than a frozen range.
    for index, offset in enumerate(SENSITIVITY_OFFSETS):
        col = 3 + index
        sign = f"+{offset}" if offset > 0 else str(offset)
        formula = f"=$B${R_TERMINAL}" if offset == 0 else f"=$B${R_TERMINAL}{sign}"
        _style(_formula(ws, R_SENS_COLS, col, formula, PERCENT_1DP, bold=True),
               bold=True, fmt=PERCENT_1DP, align="center", fill=HEADER_FILL)

    for index, offset in enumerate(SENSITIVITY_OFFSETS):
        row = R_SENS_FIRST + index
        sign = f"+{offset}" if offset > 0 else str(offset)
        formula = f"=$B${R_WACC}" if offset == 0 else f"=$B${R_WACC}{sign}"
        _style(_formula(ws, row, 2, formula, PERCENT_2DP, bold=True),
               bold=True, fmt=PERCENT_2DP, fill=HEADER_FILL)

        for col_index in range(len(SENSITIVITY_OFFSETS)):
            col = 3 + col_index
            g_ref = f"{get_column_letter(col)}${R_SENS_COLS}"
            w_ref = f"$B{row}"
            # Re-runs the whole valuation for this (WACC, g) pair. The UFCF
            # line does not depend on either, so it is reused directly; only
            # the discounting and terminal value are recomputed.
            cell_formula = (
                f'=IF({w_ref}<={g_ref},"n/a",'
                f"(SUMPRODUCT(${first}${R_UFCF}:${last}${R_UFCF},"
                f"1/(1+{w_ref})^${first}${R_PERIOD}:${last}${R_PERIOD})"
                f"+${last}${R_UFCF}*(1+{g_ref})/({w_ref}-{g_ref})"
                f"/(1+{w_ref})^${last}${R_PERIOD}"
                f"+$B${R_CASH_IN}-$B${R_DEBT_IN})/$B${R_SHARES_IN})"
            )
            cell = _formula(ws, row, col, cell_formula, PER_SHARE)
            if offset == 0 and col_index == 2:
                cell.font = Font(name=FONT_NAME, bold=True, size=10)
                cell.fill = PatternFill("solid", start_color="E2EFDA")

    # --- Notes -------------------------------------------------------------
    _label(ws, R_NOTES_HEAD, "NOTES", bold=True)
    notes = [
        "Unlevered FCF = NOPAT + D&A - capex - change in NWC.",
        "Terminal value uses Gordon growth: TV = FCF_final x (1+g) / (WACC - g), "
        "discounted at the final-year factor.",
        "Equity value = enterprise value + cash & investments - total debt.",
        "The highlighted centre cell of the sensitivity grid is the headline valuation.",
        "Blue figures are inputs; every black figure is a live formula.",
    ]
    note = honesty_note(report)
    if note:
        notes.append(note)
    for warning in report.suitability.warnings:
        notes.append(f"Warning: {warning}")
    notes.append(
        f"Generated {today} from Financial Modeling Prep data. Assumption sources "
        "are shown beside each input in section 1."
    )
    for index, text in enumerate(notes):
        _style(ws.cell(row=R_NOTES_HEAD + 1 + index, column=1, value=f"-  {text}"),
               color=GREY, size=9)

    # --- Presentation ------------------------------------------------------
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    for col in range(3, layout.last_col + 2):
        ws.column_dimensions[get_column_letter(col)].width = 15

    # No frozen panes. An earlier version froze at B24, which pinned the whole
    # assumptions block and projection header - 23 rows - above the split. On a
    # normal screen the frozen region filled most of the window, leaving a
    # scrollable strip only a few rows tall, so the sheet appeared not to
    # scroll at all until the user turned freezing off by hand.
    #
    # This model is read top to bottom in five short sections rather than as a
    # long table, and it is only ~7 columns wide, so there is nothing a freeze
    # usefully keeps in view. Left explicitly None rather than simply omitted,
    # so the intent is not mistaken for an oversight.
    ws.freeze_panes = None
    ws.sheet_view.showGridLines = False

    return wb


def workbook_bytes(report: ValuationReport) -> bytes:
    """Serialise the workbook for *report* to bytes for an HTTP response."""
    buffer = BytesIO()
    build_workbook(report).save(buffer)
    return buffer.getvalue()


def filename_for(report: ValuationReport) -> str:
    return f"{report.ticker}_DCF_Model.xlsx"
