"""
Tests for the .xlsx export.

The central claim being tested is that the workbook is *live*: assumption and
base-year cells hold numbers, and everything downstream is a formula wired to
them. A workbook of precomputed values would satisfy a naive "file opens"
check while being inert, so these tests assert formula strings and their
references explicitly.

Yahoo is mocked throughout; no network calls, no rate limits.

Run with:  python -m pytest test_excel_export.py -v
"""

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import analysis
import assumptions as A
import excel_export as E
from analysis import value_company
from api import app
from test_api import (LOSSMAKING, PROFITABLE, THROTTLED, UNKNOWN,
                      fake_fetch_financials)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def _mock_market_data(monkeypatch):
    monkeypatch.setattr(analysis, "fetch_financials", fake_fetch_financials)
    monkeypatch.setattr(A, "fetch_risk_free_rate",
                        lambda tenor="year10": (0.045, "pinned for test"))


@pytest.fixture
def report():
    return value_company(PROFITABLE)


@pytest.fixture
def sheet(report):
    """The generated workbook, round-tripped through bytes as a client sees it."""
    wb = load_workbook(BytesIO(E.workbook_bytes(report)))
    return wb["DCF Model"]


@pytest.fixture
def client():
    return TestClient(app)


def formula(sheet, row, col=2) -> str:
    value = sheet.cell(row=row, column=col).value
    return value if isinstance(value, str) else ""


# ---------------------------------------------------------------------------
# The file opens and has the expected structure
# ---------------------------------------------------------------------------

def test_workbook_opens_and_has_named_sheet(report):
    wb = load_workbook(BytesIO(E.workbook_bytes(report)))
    assert wb.sheetnames == ["DCF Model"]


def test_file_is_a_real_xlsx_zip(report):
    assert E.workbook_bytes(report)[:2] == b"PK"


def test_all_five_sections_present(sheet):
    headings = [sheet.cell(row=r, column=1).value or "" for r in range(1, 60)]
    joined = " ".join(headings)
    for section in ("1.  ASSUMPTIONS", "2.  BASE-YEAR DATA", "3.  UNLEVERED FREE CASH FLOW",
                    "4.  ENTERPRISE & EQUITY VALUE", "5.  SENSITIVITY"):
        assert section in joined, f"missing section: {section}"


def test_sheet_scrolls_freely(sheet):
    """
    Regression: freezing at B24 pinned 23 rows, which filled the viewport and
    left the sheet effectively unscrollable until the user disabled freezing.
    """
    assert sheet.freeze_panes is None


def test_downloaded_sheet_scrolls_freely(client):
    """The same guarantee through the endpoint, not just the builder."""
    r = client.get(f"/valuation/{PROFITABLE}/excel")
    assert load_workbook(BytesIO(r.content))["DCF Model"].freeze_panes is None


def test_recalculation_on_open_is_forced(report):
    """openpyxl writes no cached values, so the file must ask Excel to compute."""
    wb = load_workbook(BytesIO(E.workbook_bytes(report)))
    assert wb.calculation.fullCalcOnLoad is True


# ---------------------------------------------------------------------------
# Inputs are values; everything else is a formula
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("row,name", [
    (E.R_GROWTH, "revenue_growth"), (E.R_MARGIN, "operating_margin"),
    (E.R_TAX, "tax_rate"), (E.R_DA, "da_pct"), (E.R_CAPEX, "capex_pct"),
    (E.R_NWC, "nwc_pct"), (E.R_WACC, "wacc"), (E.R_TERMINAL, "terminal_growth"),
])
def test_assumption_cells_hold_editable_numbers(sheet, report, row, name):
    value = sheet.cell(row=row, column=2).value
    assert isinstance(value, (int, float)), f"{name} must be an editable number"
    assert value == pytest.approx(getattr(report.derived.assumptions, name))


def test_assumption_cells_are_blue_and_highlighted(sheet):
    cell = sheet.cell(row=E.R_WACC, column=2)
    assert cell.font.color.rgb.endswith("0000FF"), "inputs must be blue"
    assert cell.fill.fill_type == "solid"


def test_assumption_provenance_is_recorded_beside_each_input(sheet, report):
    source = sheet.cell(row=E.R_WACC, column=3).value
    assert source == report.derived.provenance["wacc"].source


def test_base_year_cells_hold_reported_numbers(sheet, report):
    assert sheet.cell(row=E.R_REVENUE_IN, column=2).value == pytest.approx(report.base.revenue)
    assert sheet.cell(row=E.R_SHARES_IN, column=2).value == pytest.approx(report.base.shares)
    assert sheet.cell(row=E.R_PRICE_IN, column=2).value == pytest.approx(report.base.current_price)


# ---------------------------------------------------------------------------
# Projection formulas
# ---------------------------------------------------------------------------

def test_revenue_compounds_by_the_growth_assumption(sheet):
    assert formula(sheet, E.R_REVENUE, 3) == f"=B{E.R_REVENUE}*(1+$B${E.R_GROWTH})"
    assert formula(sheet, E.R_REVENUE, 4) == f"=C{E.R_REVENUE}*(1+$B${E.R_GROWTH})"


def test_base_revenue_links_to_the_input_cell(sheet):
    assert formula(sheet, E.R_REVENUE, 2) == f"=$B${E.R_REVENUE_IN}"


@pytest.mark.parametrize("row,reference", [
    (E.R_EBIT, "$B$6"), (E.R_TAXES, "$B$7"), (E.R_DA_ADD, "$B$8"),
    (E.R_CAPEX_LESS, "$B$9"), (E.R_NWC_LESS, "$B$10"),
])
def test_ufcf_lines_reference_their_assumption(sheet, row, reference):
    text = formula(sheet, row, 3)
    assert text.startswith("="), "must be a formula, not a value"
    assert reference in text, f"row {row} should be driven by {reference}"


def test_discount_factor_uses_wacc_and_period_row(sheet):
    text = formula(sheet, E.R_DISCOUNT, 3)
    assert text == f"=1/(1+$B${E.R_WACC})^C${E.R_PERIOD}"


def test_every_projection_cell_is_a_formula(sheet, report):
    rows = [E.R_REVENUE, E.R_EBIT, E.R_TAXES, E.R_NOPAT, E.R_DA_ADD,
            E.R_CAPEX_LESS, E.R_NWC_LESS, E.R_UFCF, E.R_DISCOUNT, E.R_PV_UFCF]
    last_col = 2 + report.derived.assumptions.projection_years
    for row in rows:
        for col in range(3, last_col + 1):
            assert formula(sheet, row, col).startswith("="), \
                f"cell R{row}C{col} is not a formula"


# ---------------------------------------------------------------------------
# Terminal value and equity bridge
# ---------------------------------------------------------------------------

def test_terminal_value_is_a_gordon_growth_formula(sheet):
    text = formula(sheet, E.R_TV)
    assert text.startswith("=")
    assert f"(1+$B${E.R_TERMINAL})" in text
    assert f"($B${E.R_WACC}-$B${E.R_TERMINAL})" in text


def test_equity_bridge_adds_cash_and_subtracts_debt(sheet):
    assert formula(sheet, E.R_PLUS_CASH) == f"=$B${E.R_CASH_IN}"
    assert formula(sheet, E.R_LESS_DEBT) == f"=-$B${E.R_DEBT_IN}"
    assert formula(sheet, E.R_EQUITY) == \
        f"=$B${E.R_EV}+$B${E.R_PLUS_CASH}+$B${E.R_LESS_DEBT}"


def test_intrinsic_value_is_a_formula_not_a_number(sheet):
    text = formula(sheet, E.R_INTRINSIC)
    assert text == f"=$B${E.R_EQUITY}/$B${E.R_SHARES}", \
        "the headline figure must be computed by Excel, not baked in"


def test_pv_sum_spans_the_projection(sheet, report):
    last = chr(ord("C") + report.derived.assumptions.projection_years - 1)
    assert formula(sheet, E.R_PV_SUM) == f"=SUM(C{E.R_PV_UFCF}:{last}{E.R_PV_UFCF})"


# ---------------------------------------------------------------------------
# Sensitivity grid
# ---------------------------------------------------------------------------

def test_sensitivity_axes_are_anchored_to_the_assumption_cells(sheet):
    """So the grid re-centres when the user edits WACC or g."""
    assert formula(sheet, E.R_SENS_COLS, 3) == f"=$B${E.R_TERMINAL}-0.01"
    assert formula(sheet, E.R_SENS_COLS, 5) == f"=$B${E.R_TERMINAL}"
    assert formula(sheet, E.R_SENS_FIRST, 2) == f"=$B${E.R_WACC}-0.01"
    assert formula(sheet, E.R_SENS_FIRST + 2, 2) == f"=$B${E.R_WACC}"


def test_all_25_sensitivity_cells_are_formulas(sheet):
    for row in range(E.R_SENS_FIRST, E.R_SENS_LAST + 1):
        for col in range(3, 8):
            text = formula(sheet, row, col)
            assert text.startswith("=SUMPRODUCT") or text.startswith("=IF"), \
                f"sensitivity cell R{row}C{col} is not a live formula"


def test_sensitivity_cells_recompute_the_whole_valuation(sheet):
    text = formula(sheet, E.R_SENS_FIRST, 3)
    assert "SUMPRODUCT" in text, "must re-discount the UFCF line"
    assert f"$B${E.R_CASH_IN}" in text and f"$B${E.R_DEBT_IN}" in text, "must apply the equity bridge"
    assert f"$B${E.R_SHARES_IN}" in text, "must divide by shares"


def test_sensitivity_guards_against_growth_exceeding_wacc(sheet):
    text = formula(sheet, E.R_SENS_FIRST, 3)
    assert text.startswith("=IF(") and '"n/a"' in text


def test_no_sensitivity_cell_contains_a_hardcoded_number(sheet):
    for row in range(E.R_SENS_FIRST, E.R_SENS_LAST + 1):
        for col in range(3, 8):
            assert not isinstance(sheet.cell(row=row, column=col).value, (int, float))


# ---------------------------------------------------------------------------
# Refusals
# ---------------------------------------------------------------------------

def test_building_a_workbook_for_a_refused_company_raises():
    refused = value_company(LOSSMAKING)
    with pytest.raises(ValueError, match="suitability guard"):
        E.build_workbook(refused)


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

def test_excel_endpoint_returns_a_downloadable_file(client):
    r = client.get(f"/valuation/{PROFITABLE}/excel")
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX_MEDIA_TYPE
    assert f'filename="{PROFITABLE}_DCF_Model.xlsx"' in r.headers["content-disposition"]
    assert r.content[:2] == b"PK"


def test_downloaded_file_opens_and_is_live(client):
    r = client.get(f"/valuation/{PROFITABLE}/excel")
    ws = load_workbook(BytesIO(r.content))["DCF Model"]
    assert str(ws.cell(row=E.R_INTRINSIC, column=2).value).startswith("=")


def test_excel_endpoint_refuses_unsuitable_company(client):
    r = client.get(f"/valuation/{LOSSMAKING}/excel")
    assert r.status_code == 422
    assert r.json()["code"] == "not_suitable"
    assert r.headers["content-type"].startswith("application/json"), \
        "a refusal must not masquerade as a spreadsheet"


def test_excel_endpoint_reports_upstream_rate_limiting(client):
    r = client.get(f"/valuation/{THROTTLED}/excel")
    assert r.status_code == 429
    assert r.json()["code"] == "upstream_rate_limited"


def test_excel_endpoint_reports_unknown_ticker(client):
    r = client.get(f"/valuation/{UNKNOWN}/excel")
    assert r.status_code == 404
    assert r.json()["code"] == "ticker_not_found"


def test_excel_post_bakes_overrides_into_the_input_cells(client):
    r = client.post("/valuation/excel",
                    json={"ticker": PROFITABLE, "overrides": {"wacc": 0.123}})
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content))["DCF Model"]
    assert ws.cell(row=E.R_WACC, column=2).value == pytest.approx(0.123)
    assert ws.cell(row=E.R_WACC, column=3).value == "override"


def test_excel_endpoints_appear_in_the_openapi_schema(client):
    paths = client.get("/openapi.json").json()["paths"]
    assert "/valuation/{ticker}/excel" in paths
    assert XLSX_MEDIA_TYPE in paths["/valuation/{ticker}/excel"]["get"]["responses"]["200"]["content"]
