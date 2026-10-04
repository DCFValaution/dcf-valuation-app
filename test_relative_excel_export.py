"""
Tests for the relative (peer multiples) workbook.

This one is a table rather than a model, so the claims are different: that the
comparison is all there - every peer, every multiple, the medians - that the
arithmetic a reader might argue with is live, that the peer group's provenance
travels with it, and that the market-based framing is in the sheet rather than
left on the screen it came from.

The guard matters as much as the content: where the relative view declined,
there is no figure, and the builder refuses rather than producing a sheet laid
out like a valuation with the number missing.

Run with:  python -m pytest test_relative_excel_export.py -v
"""

import pytest
from openpyxl import load_workbook
from io import BytesIO

import relative_excel_export as X
from test_api import PROFITABLE as COMPANY
from test_relative_api import _stubbed, client  # noqa: F401  (fixtures)
import analysis


@pytest.fixture
def report(client):  # noqa: F811
    return analysis.value_company_relatively(COMPANY)


@pytest.fixture
def sheet(report):
    return X.build_relative_workbook(report)["Relative view"]


def _column_a(ws, limit=90):
    return [str(ws.cell(row=r, column=1).value or "") for r in range(1, limit)]


def _text(ws, limit=90):
    return " ".join(_column_a(ws, limit)).lower()


def _find(ws, needle, limit=90):
    for row in range(1, limit):
        value = ws.cell(row=row, column=1).value
        if value and needle.lower() in str(value).lower():
            return row
    return None


def _header_row(ws, limit=90):
    """The table's column headers, not the section title above them.

    "1.  MULTIPLES   (this company beside its peers)" contains the word too,
    so a substring search finds the banner and reads the wrong row.
    """
    for row in range(1, limit):
        if str(ws.cell(row=row, column=1).value or "").strip() == "Multiple":
            return row
    raise AssertionError("no column-header row found")


# ---------------------------------------------------------------------------
# It is a comparison, and the comparison is all there
# ---------------------------------------------------------------------------

def test_every_peer_has_a_column(sheet, report):
    header = _header_row(sheet)
    headers = [str(sheet.cell(row=header, column=col).value or "")
               for col in range(2, 12)]
    for peer in report.peers:
        assert peer.ticker in headers, f"{peer.ticker} has no column"


def test_the_company_sits_beside_its_peers(sheet, report):
    header = _header_row(sheet)
    assert str(sheet.cell(row=header, column=2).value) == report.ticker


def test_every_multiple_has_a_row(sheet, report):
    labels = _text(sheet)
    for multiple in report.result.multiples:
        assert multiple.label.lower() in labels


def test_the_peer_group_is_named_with_its_provenance(sheet, report):
    """Who the peers are, and how each got there - because changing the group
    changes the figure."""
    assert _find(sheet, "THE PEER GROUP") is not None
    text = " ".join(
        str(sheet.cell(row=r, column=c).value or "")
        for r in range(1, 90) for c in range(1, 5)).lower()
    for peer in report.peers:
        assert peer.ticker.lower() in text
        assert peer.provenance.lower() in text


# ---------------------------------------------------------------------------
# The arithmetic a reader might argue with is live
# ---------------------------------------------------------------------------

def test_the_medians_are_formulas_over_the_peer_cells(sheet, report):
    header = _header_row(sheet)
    median_col = 3 + len(report.peers)
    for offset, multiple in enumerate(report.result.multiples):
        if multiple.peer_median is None:
            continue
        cell = sheet.cell(row=header + 1 + offset, column=median_col)
        assert isinstance(cell.value, str) and cell.value.startswith("=MEDIAN(")


def test_the_premium_to_the_median_is_a_formula(sheet, report):
    row = _find(sheet, "premium / (discount)")
    assert row is not None
    assert str(sheet.cell(row=row, column=2).value).startswith("=")


def test_the_central_figure_is_a_formula_over_the_implied_values(sheet):
    row = _find(sheet, "Central value per share")
    assert row is not None
    assert str(sheet.cell(row=row, column=2).value).startswith("=MEDIAN(")


# ---------------------------------------------------------------------------
# The framing
# ---------------------------------------------------------------------------

def test_the_sheet_says_it_is_not_intrinsic_value(sheet):
    assert "not an estimate of intrinsic value" in _text(sheet)


def test_the_mispricing_caveat_travels_with_the_file(sheet):
    text = _text(sheet)
    assert "inherits the market" in text
    assert "mispricing" in text


def test_it_says_the_peer_group_is_a_judgement(sheet):
    assert "change the peer group and the figure changes" in _text(sheet)


# ---------------------------------------------------------------------------
# The guard
# ---------------------------------------------------------------------------

def test_a_declined_comparison_produces_no_workbook(report, monkeypatch):
    """No figure on screen means no figure in a file.

    A sheet laid out like a valuation with one empty cell invites the reader
    to fill it in, which is worse than refusing.
    """
    declined = report
    monkeypatch.setattr(type(declined.suitability), "suitable",
                        property(lambda self: False), raising=False)
    with pytest.raises(ValueError, match="Refusing to build"):
        X.build_relative_workbook(declined)


def test_it_refuses_a_report_of_another_kind(client):  # noqa: F811
    plain = analysis.value_company(COMPANY)
    with pytest.raises(ValueError, match="this report is a"):
        X.build_relative_workbook(plain)


# ---------------------------------------------------------------------------
# Housekeeping
# ---------------------------------------------------------------------------

def test_it_recalculates_on_open(report):
    buffer = BytesIO(X.relative_workbook_bytes(report))
    assert load_workbook(buffer).calculation.fullCalcOnLoad is True


def test_nothing_is_frozen(sheet):
    assert sheet.freeze_panes is None


def test_the_filename_names_the_view(report):
    assert X.relative_filename_for(report) == f"{report.ticker}_Relative.xlsx"
